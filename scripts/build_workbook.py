"""Build the Excel workbook from the cleaned dataset. India-only sources, last-12-months data."""
import re
import sys
from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
B = r"C:\Users\RISHABH\Desktop\Pediatric_Medicine_India"
RAW = B + r"\raw_pages"
OUT = B + r"\India_Pediatric_Liquid_Medicines_Oct2026.xlsx"
TOP_FAMILIES = 30

df = pd.read_csv(RAW + r"\dataset_clean.csv", encoding="utf-8-sig")
excl = pd.read_csv(RAW + r"\dataset_excluded.csv", encoding="utf-8-sig")
ev = pd.read_csv(RAW + r"\evidence_table.csv", encoding="utf-8-sig")
ph = pd.read_csv(RAW + r"\pharmarack_brand_rows.csv", encoding="utf-8-sig")

BOBCAPS_URL = "https://www.barodaetrade.com/Reports/Pharmaceuticals-SectorUpdate13Aug26-Research.pdf"
BIOSPECTRUM_URL = "https://biospectrumindia.com/news/73/27057/mounjaro-emerges-as-top-brand-of-november-as-pharmarack-reports-robust-india-pharma-market-growth.html"
NLEM_READ = "https://www.cspc.in/pdf2023/Essential_Drugs_National.pdf"
NLEM_OFFICIAL = "https://www.cdsco.gov.in/opencms/resources/UploadCDSCOWeb/2018/UploadPublic_NoticesFiles/reportnlem2022.pdf"
NLEM_CURRENT_URL = "https://superkalam.com/current-affairs/05-07-2026/india-s-essential-medicine-list-lags-behind-global-benchmark-unrevised-in-4-years-pg1-c2f6c8d7-a1b7-4370-9706-08fc7025ee82"
REGION_TXT = "Not publicly available. Platforms list products nationally; state/city sales are only in paid IQVIA/Pharmarack zone reports."

# Pharmarack brand -> (strict brand family in this dataset, extra regex the product name must also match, or None)
PHARM_MAP = {
    "CALPOL": ("Calpol", None), "AUGMENTIN": ("Augmentin Duo", None), "CLAVAM": ("Clavam", None), "TAXIM O": ("Taxim-O", None), "ONDEM": ("Ondem", None), "ZOCEF": ("Zocef", None),
    "ECONORM": ("Econorm", None), "AZEE": ("Azee", None), "ASTHALIN": ("Asthalin", r"^(?!.*\bdx\b)"), "MONTAIR LC": ("Montair", r"\blc\b"), "DUPHALAC": ("Duphalac", None),
    "CREMAFFIN PLUS": ("Cremaffin", r"\bplus\b"), "COREX DX": ("Corex", r"\bdx\b"), "CEFTUM": ("Ceftum", None),
}
phd = {r["brand"]: r for r in ph.to_dict("records")}


def pharm_for(name, family):
    n = str(name).lower()
    for brand, (fam_name, extra) in PHARM_MAP.items():
        if family == fam_name and brand in phd and (extra is None or re.search(extra, n)):
            return phd[brand]
    return None


pj, pjl, ppg, pbr = [], [], [], []
for n, famly in zip(df["name"], df["brand_family"]):
    r = pharm_for(n, famly)
    pbr.append(r["brand"] if r else None)
    pj.append(r["sv_jun_2026_rs_mn"] if r else None)
    pjl.append(r["sv_jul_2026_rs_mn"] if r else None)
    ppg.append(int(r["page"]) if r else None)
df["pharm_brand"], df["pharm_jun"], df["pharm_jul"], df["pharm_page"] = pbr, pj, pjl, ppg

# ---------------------------------------------------------------- family-level evidence tier
fam = df.groupby("brand_family").agg(pharm_jul=("pharm_jul", "max"), bought_max=("bought_n", "max"), rating_max=("rating_count", "max"), n_sku=("url", "count")).reset_index()


def tier(r):
    if pd.notna(r["pharm_jul"]):
        return "A"
    if pd.notna(r["bought_max"]):
        return "B"
    if pd.notna(r["rating_max"]):
        return "C"
    return None


fam["tier"] = fam.apply(tier, axis=1)
fam["value"] = fam.apply(lambda r: r["pharm_jul"] if r["tier"] == "A" else (r["bought_max"] if r["tier"] == "B" else (r["rating_max"] if r["tier"] == "C" else None)), axis=1)
fam_ev = fam[fam["tier"].notna()].sort_values(["tier", "value"], ascending=[True, False]).reset_index(drop=True)
fam_ev["family_rank"] = fam_ev.index + 1
top = fam_ev.head(TOP_FAMILIES)
TIER_TXT = {
    "A": "A - National sales figure (Pharmarack, Jun/Jul 2026)",
    "B": "B - Platform purchase count (1mg 'people bought recently')",
    "C": "C - Platform ratings count",
}
df = df.merge(fam[["brand_family", "tier", "value"]].rename(columns={"tier": "fam_tier", "value": "fam_value"}), on="brand_family", how="left")
df = df.merge(fam_ev[["brand_family", "family_rank"]], on="brand_family", how="left")

# ---------------------------------------------------------------- text columns
def sources_text(r):
    parts = []
    if pd.notna(r["pharm_jul"]):
        parts.append(f"BOBCAPS Securities, Pharmaceuticals Sector Update, 13-Aug-2026, p{int(r['pharm_page'])}, data: Pharmarack - {BOBCAPS_URL}")
    if r["brand_family"] == "Augmentin Duo":
        parts.append(f"Pharmarack MAT ranking, BioSpectrum India, 11-Dec-2025 - {BIOSPECTRUM_URL}")
    return " | ".join(parts) if parts else "No public national sales figure found for this brand in the last 12 months"


def national_text(r):
    if pd.notna(r["pharm_jul"]):
        t = f"{r['pharm_brand']}: Rs {r['pharm_jun']:.0f} mn (Jun-2026), Rs {r['pharm_jul']:.0f} mn (Jul-2026) - BRAND level, all dosage forms"
        if r["brand_family"] == "Augmentin Duo":
            t += "; ranked #1 India brand at MAT level (Pharmarack, Dec-2025)"
        return t
    return "No public national sales figure found"


def platform_signal(r):
    return r["demand_text"] if isinstance(r["demand_text"], str) and r["demand_text"] else "No demand figure shown on this platform page"


def yesno(v):
    if v is True or str(v).lower() == "true":
        return "Yes"
    if v is False or str(v).lower() == "false":
        return "No"
    return "Not shown"


df["national_txt"] = df.apply(national_text, axis=1)
df["national_src"] = df.apply(sources_text, axis=1)
df["platform_signal"] = df.apply(platform_signal, axis=1)
df["in_stock"] = df["available"].apply(yesno)
df["retrieved"] = df["retrieved_utc"].str[:10]
df["tier_txt"] = df["fam_tier"].map(TIER_TXT).fillna("None - no public demand figure found")
def mfr_text(r):
    m, k = r["manufacturer"], r["marketer"]
    m = m if isinstance(m, str) and m.strip() else None
    k = k if isinstance(k, str) and k.strip() else None
    if m and k and m != k:
        return f"Manufacturer: {m}; Marketer: {k}"
    return m or k or "Not shown"


df["mfr"] = df.apply(mfr_text, axis=1)
df["location_txt"] = df["location"].apply(lambda v: v if isinstance(v, str) and v else "Not shown on this page")
df["region_txt"] = REGION_TXT
df["nlem_src"] = f"NLEM 2022 (Govt of India, 13-Sep-2022; still the current list as of Jul-2026). Read: {NLEM_READ} | Official: {NLEM_OFFICIAL}"

# ---------------------------------------------------------------- sheet 1
s1 = df[df["family_rank"].isin(top["family_rank"])].copy()
s1["_bn"] = s1["bought_n"].fillna(-1)
s1 = s1.sort_values(["family_rank", "_bn", "site", "mrp"], ascending=[True, False, True, True])
S1_COLS = [
    ("Rank (brand)", "family_rank"), ("Brand demand-evidence tier", "tier_txt"), ("Brand family", "brand_family"), ("Product (as listed)", "name"), ("Platform", "site"),
    ("Molecule + strength (as listed on page)", "composition"), ("Formulation", "formulation"), ("Pediatric label in product name", "pediatric_label"),
    ("Manufacturer / marketer", "mfr"), ("MRP (Rs)", "mrp"), ("Pack size (as listed)", "pack"), ("Other pack sizes & MRP (1mg)", "other_packs"), ("In stock on page", "in_stock"),
    ("Platform demand signal (verbatim)", "platform_signal"), ("National sales evidence (Pharmarack)", "national_txt"), ("National evidence source", "national_src"),
    ("Demand by state / city", "region_txt"), ("Manufacturer / marketer location (from page)", "location_txt"), ("Product page URL (source)", "url"), ("Data retrieved", "retrieved"),
]
sheet1 = s1[[c for _, c in S1_COLS]].rename(columns={c: h for h, c in S1_COLS}).fillna("Not shown")

# ---------------------------------------------------------------- sheet 2
s2 = df[df["rel_level"] <= 4].copy()
s2["_t"] = s2["fam_tier"].map({"A": 0, "B": 1, "C": 2}).fillna(3)
s2["_v"] = -s2["fam_value"].fillna(0)
s2["_p"] = (s2["pediatric_label"] != "Yes").astype(int)
s2 = s2.sort_values(["rel_level", "_t", "_v", "_p", "molecule_key", "site", "mrp"]).reset_index(drop=True)
s2.insert(0, "relevance_rank", s2.index + 1)
S2_COLS = [
    ("Relevance rank", "relevance_rank"), ("Relevance level (India NLEM 2022)", "rel_level"), ("Relevance basis", "rel_basis"), ("Brand family", "brand_family"),
    ("Product (as listed)", "name"), ("Platform", "site"), ("Molecule group", "molecule_key"), ("Molecule + strength (as listed on page)", "composition"), ("Formulation", "formulation"),
    ("Pediatric label in product name", "pediatric_label"), ("Manufacturer / marketer", "mfr"), ("MRP (Rs)", "mrp"), ("Pack size (as listed)", "pack"),
    ("NLEM 2022 list row (verbatim extract)", "nlem_row"), ("NLEM location in document", "nlem_pages"), ("NLEM source", "nlem_src"),
    ("Brand demand-evidence tier", "tier_txt"), ("Platform demand signal (verbatim)", "platform_signal"), ("National sales evidence (Pharmarack)", "national_txt"),
    ("Demand by state / city", "region_txt"), ("Manufacturer / marketer location (from page)", "location_txt"), ("Product page URL (source)", "url"), ("Data retrieved", "retrieved"),
]
sheet2 = s2[[c for _, c in S2_COLS]].rename(columns={c: h for h, c in S2_COLS}).fillna("Not shown")

# ---------------------------------------------------------------- supporting sheets
lvl_of = {}
for mol_key, (lvl, ekey) in {
    "Paracetamol": (1, "Paracetamol"), "Ibuprofen": (1, "Ibuprofen"), "Mefenamic acid": (1, "Mefenamic acid"), "Amoxicillin": (1, "Amoxicillin"),
    "Amoxicillin + Clavulanic acid": (1, "Amoxicillin + Clavulanic acid"), "Azithromycin": (1, "Azithromycin"), "Cefuroxime": (1, "Cefuroxime"), "Cetirizine": (1, "Cetirizine"),
    "Albendazole": (1, "Albendazole"), "Cefadroxil": (1, "Cefadroxil"), "Co-trimoxazole": (1, "Co-trimoxazole"), "Metronidazole": (1, "Metronidazole"), "Prednisolone": (1, "Prednisolone"),
    "Ferrous salt (+folic acid)": (1, "Ferrous salt (+folic acid)"), "Cefixime": (2, "Cefixime"), "Ondansetron": (2, "Ondansetron"), "Clarithromycin": (2, "Clarithromycin"),
    "Domperidone": (3, "Domperidone"), "Lactulose": (3, "Lactulose"), "Salbutamol": (3, "Salbutamol"), "Oral rehydration salts": (3, "Oral rehydration salts"), "Chlorpheniramine": (3, "Chlorpheniramine"),
    "Zinc sulphate": (4, "Zinc sulphate"), "Montelukast": (4, "Montelukast"),
}.items():
    lvl_of[ekey] = lvl
ev["Relevance level"] = ev["molecule"].map(lvl_of)
ev_out = ev.rename(columns={"molecule": "Molecule", "nlem_2022_row": "NLEM 2022 list row (verbatim extract)", "nlem_pdf_page": "PDF page", "nlem_printed_page": "Printed page", "nlem_doc_url_read": "Document read (URL)", "nlem_official_url": "Official source (URL)"})
ph_out = ph.rename(columns={"page": "PDF page", "company": "Company", "rank_in_company": "Rank within company", "brand": "Brand", "sv_jun_2026_rs_mn": "Sales value Jun-2026 (Rs mn)", "sv_jul_2026_rs_mn": "Sales value Jul-2026 (Rs mn)"}).copy()
ph_out["Source"] = f"BOBCAPS Securities, Pharmaceuticals Sector Update, 13-Aug-2026 (data: Pharmarack). Brand-level, ALL dosage forms. {BOBCAPS_URL}"

allcols = ["site", "brand_family", "name", "molecule_key", "composition", "formulation", "pediatric_label", "manufacturer", "marketer", "mrp", "pack", "other_packs", "in_stock",
           "platform_signal", "national_txt", "rel_level", "rel_basis", "url", "retrieved"]
all_out = df[allcols].sort_values(["site", "brand_family", "name"]).rename(columns={"site": "Platform", "brand_family": "Brand family", "name": "Product", "molecule_key": "Molecule group",
                                                                              "composition": "Composition (as listed)", "formulation": "Formulation", "pediatric_label": "Pediatric label in name",
                                                                              "manufacturer": "Manufacturer", "marketer": "Marketer", "mrp": "MRP (Rs)", "pack": "Pack size", "other_packs": "Other packs & MRP",
                                                                              "in_stock": "In stock", "platform_signal": "Platform demand signal", "national_txt": "National sales evidence",
                                                                              "rel_level": "NLEM relevance level", "rel_basis": "Relevance basis", "url": "Product page URL", "retrieved": "Retrieved"})
all_out = all_out.fillna("Not shown")
excl_out = excl.rename(columns={"site": "Platform", "url": "URL", "name": "Product", "reason": "Why excluded"})

cov = df.groupby("site").agg(products=("url", "count"), with_mrp=("mrp", lambda s: s.notna().sum()), with_demand=("platform_signal", lambda s: (s != "No demand figure shown on this platform page").sum())).reset_index()
cov_rows = [
    ["Platform", "Products kept", "With MRP", "With a platform demand figure", "Notes"],
]
NOTES = {
    "1mg": "Shows 'N people bought recently' on some products (demand signal) and ratings on a few.",
    "PharmEasy": "Shows price, pack, manufacturer, molecule. No purchase or ratings counts found on these products.",
    "Apollo": "Shows price, manufacturer, dosage form. No demand figure on the page. Pack size taken from the product name.",
    "Netmeds": "159 pages fetched (capped at 3 per brand family because robots.txt sets a 10 s crawl delay). 136 were dropped because the listing's price-valid-until date had already passed (stale), leaving the rows shown. Ratings rarely present.",
}
for r in cov.to_dict("records"):
    cov_rows.append([r["site"], r["products"], r["with_mrp"], r["with_demand"], NOTES.get(r["site"], "")])
cov_rows += [
    ["MedPlus", 0, 0, 0, "NOT COVERED. medplusmart.com returns HTTP 403 (blocked) to automated requests, and I did not work around the block."],
    [],
    ["Not found / not available", "", "", "", ""],
    ["Demand by state/city", "", "", "", "No public source. Platforms show national listings; state/city sales exist only in paid IQVIA/Pharmarack zone data."],
    ["Pediatric-specific sales ranking", "", "", "", "No public ranking of pediatric syrups exists. Pharmarack brand figures are brand-level across ALL dosage forms (tablets, injections included)."],
    ["Age band 2-12 years per SKU", "", "", "", "Not verified per product. Relevance is shown at molecule/dosage-form level (NLEM 2022) plus a 'pediatric label in name' flag."],
    ["NLEM '(p)' tag", "", "", "", "Appears on oral-liquid forms in the list; its legend was not found in the extracted document text, so it is quoted verbatim and not interpreted beyond that."],
]

README = [
    ["India pediatric liquid medicines - 1mg, PharmEasy, Apollo, Netmeds (retrieved 31 Oct 2026)"],
    [""],
    ["Scope", "India only. Data from the last 12 months (since 31 Oct 2025). Formulations: syrup/oral solution, suspension, dry powder for suspension (dry syrup), sachet, granules."],
    ["Platforms", "1mg, PharmEasy, Apollo Pharmacy, Netmeds. MedPlus is NOT covered (it blocks automated access; I did not work around it)."],
    ["How products were found", "Each site's public product sitemap (listed in its robots.txt), filtered to named, nationally known pediatric brands. Only pages allowed by robots.txt were fetched, at each site's crawl delay."],
    ["Selection caveat", "The brand list is my selection of well-known pediatric liquid brands, not a census of the Indian market. A high-selling brand missing from the list will be missing here."],
    [""],
    ["Sheet 1 - Most Sold / High Demand", f"Top {TOP_FAMILIES} brand families that have a demand figure. Ranked by evidence quality first (A, then B, then C), then by size. A = Pharmarack national sales (Rs mn, Jun/Jul 2026, brand level, all dosage forms). B = 1mg 'people bought recently' (count shown on the page; the time window is not stated by 1mg). C = ratings count."],
    ["Important", "No source publishes pediatric-only or suspension-only sales. National figures are brand totals across all forms. Platform counts are not unit sales."],
    ["Sheet 2 - Most Pediatric Relevant", "Ranked by India's National List of Essential Medicines 2022 (NLEM). Level 1 = oral liquid tagged '(p)' at primary-care level; 2 = '(p)' at secondary/tertiary level; 3 = listed without '(p)'; 4 = listed only in a different form or as separate components. Level 5 (not on NLEM) is in 'All_data'."],
    ["NLEM date note", f"NLEM 2022 was published 13 Sep 2022, which is older than 12 months. It is kept because it is still India's current official list (not revised as of Jul 2026: {NLEM_CURRENT_URL}). If you want strictly 12-month sources only, ignore the relevance columns."],
    ["Molecule labels", "Taken from each page's own listed composition. Where a page lists none, the brand family was used as a fallback."],
    ["Prices", "MRP is the price shown on the product page on the retrieval date. Pack size is as listed on that platform."],
    ["Recency filters applied", "Dropped PharmEasy pages last updated before 31 Oct 2025 and Netmeds listings whose price-valid-until date had already passed. Veterinary (pet) products were removed."],
    ["Sources", f"Platform product pages (URLs in each row). Pharmarack data via BOBCAPS Securities report: {BOBCAPS_URL} . Augmentin MAT rank: {BIOSPECTRUM_URL} . NLEM 2022: {NLEM_OFFICIAL}"],
]

with pd.ExcelWriter(OUT, engine="openpyxl") as xw:
    pd.DataFrame(README).to_excel(xw, sheet_name="README", index=False, header=False)
    sheet1.to_excel(xw, sheet_name="1_Most_Sold_High_Demand", index=False)
    sheet2.to_excel(xw, sheet_name="2_Most_Pediatric_Relevant", index=False)
    ev_out.to_excel(xw, sheet_name="NLEM_2022_evidence", index=False)
    ph_out.to_excel(xw, sheet_name="Pharmarack_brand_sales", index=False)
    all_out.to_excel(xw, sheet_name="All_data", index=False)
    pd.DataFrame(cov_rows).to_excel(xw, sheet_name="Coverage_and_gaps", index=False, header=False)
    excl_out.to_excel(xw, sheet_name="Excluded_rows", index=False)

# ---------------------------------------------------------------- formatting
wb = load_workbook(OUT)
HEAD = PatternFill("solid", fgColor="1F3864")
WIDTH = {"Product (as listed)": 38, "Product": 38, "Molecule + strength (as listed on page)": 34, "Composition (as listed)": 34, "Relevance basis": 52, "NLEM 2022 list row (verbatim extract)": 60,
         "NLEM source": 40, "National evidence source": 50, "Platform demand signal (verbatim)": 30, "National sales evidence (Pharmarack)": 38, "National sales evidence": 38,
         "Demand by state / city": 34, "Product page URL (source)": 46, "Product page URL": 46, "Other pack sizes & MRP (1mg)": 34, "Other packs & MRP": 34,
         "Manufacturer / marketer": 32, "Manufacturer / marketer location (from page)": 36, "Brand demand-evidence tier": 30, "Why excluded": 50, "URL": 46,
         "Official source (URL)": 46, "Document read (URL)": 46, "Source": 70}
for ws in wb.worksheets:
    if ws.title in ("README", "Coverage_and_gaps"):
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 120 if ws.title == "README" else 16
        for col in "CD" if ws.title == "Coverage_and_gaps" else "":
            ws.column_dimensions[col].width = 22
        if ws.title == "Coverage_and_gaps":
            ws.column_dimensions["E"].width = 110
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws["A1"].font = Font(bold=True, size=13)
        if ws.title == "Coverage_and_gaps":
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = HEAD
        else:
            for r in range(3, ws.max_row + 1):
                ws.cell(r, 1).font = Font(bold=True)
        continue
    ws.freeze_panes = "E2" if ws.title.startswith(("1_", "2_")) else "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, c in enumerate(ws[1], start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(str(c.value), 16)
    url_cols = [i for i, c in enumerate(ws[1], start=1) if "URL" in str(c.value) and "source" in str(c.value).lower() or str(c.value) in ("Product page URL", "URL", "Document read (URL)", "Official source (URL)")]
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i in url_cols:
            cell = row[i - 1]
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
wb.save(OUT)
print("saved", OUT)
print("sheet 1 rows:", len(sheet1), "| families:", s1["brand_family"].nunique())
print("sheet 2 rows:", len(sheet2))
print("all rows:", len(all_out), "| excluded:", len(excl_out))
print("\nTop families (rank, tier, value):")
print(top[["family_rank", "brand_family", "tier", "value", "n_sku"]].to_string(index=False))
